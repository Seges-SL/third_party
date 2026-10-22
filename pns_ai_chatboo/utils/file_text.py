# -*- coding: utf-8 -*-
# SPDX-License-Identifier: Apache-2.0
# Copyright (C) 2025-2026 Patanegra Soft <https://patanegra.com>
"""The text of a file the user clipped to a chat turn.

Bytes in, text out: Excel (every sheet), Word, plain text and data. No Odoo,
no session. A PDF is not read here: its pages and scans go through pns_base
and the turn's vision chain.

The text goes to the model, never to the screen, so it is in English: the
model answers in the user's language anyway. Nothing is truncated: the model
sees the whole file, every Excel sheet included.
"""

import base64
import logging

_logger = logging.getLogger(__name__)


def split_data_url(data_url):
    """Split a ``data:<mime>;base64,<b64>`` URL into ``(mimetype, b64)``.

    Bare base64 (no prefix) is taken as PNG.
    """
    s = data_url or ''
    if s.startswith('data:'):
        try:
            head, b64 = s.split(',', 1)
        except ValueError:
            return 'image/png', ''
        mimetype = head[5:].split(';', 1)[0] or 'image/png'
        return mimetype, b64
    return 'image/png', s


def decode_text_bytes(raw):
    """Bytes to text: UTF-8 (with or without BOM), then Latin-1, then replace."""
    for enc in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return raw.decode('utf-8', 'replace')


def peek_ooxml_kind(raw):
    """The real OOXML kind from the ZIP signature and its inner paths, whatever the name says.

    A ``book.xlsx`` renamed ``.xls`` (or a ``.docx`` renamed ``.doc``) is still
    read; otherwise its binary was decoded as text.
    """
    if not raw or raw[:2] != b'PK':
        return None
    import io
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            names = zf.namelist()
    except Exception:
        return None
    if any(n.startswith('xl/') for n in names):
        return 'xlsx'
    if 'word/document.xml' in names or any(n.startswith('word/') for n in names):
        return 'docx'
    if any(n.startswith('ppt/') for n in names):
        return 'pptx'
    return 'zip'


def is_ole_cfb(raw):
    """True for a Compound File Binary (old ``.xls`` / ``.doc``)."""
    return bool(raw) and raw[:4] == b'\xD0\xCF\x11\xE0'


def format_excel_sheets(sheet_rows):
    """Every sheet under its own header. No row or sheet is dropped.

    ``sheet_rows``: list of ``(title, list of tab-separated rows)``.
    """
    total = len(sheet_rows)
    titles = [t for t, _rows in sheet_rows]
    intro = (
        'Excel workbook: %d sheet(s), ALL listed (not truncated): %s'
        % (total, ', '.join('"%s"' % t for t in titles))
    )
    blocks = [intro]
    for pos, (title, rows) in enumerate(sheet_rows, start=1):
        header = '===== SHEET %d/%d: "%s" (%d row(s) with data) =====' % (
            pos, total, title, len(rows),
        )
        body = '\n'.join(rows) if rows else '(empty sheet)'
        blocks.append('%s\n%s' % (header, body))
    return '\n\n'.join(blocks)


def xlsx_to_text(raw):
    """An .xlsx/.xlsm as text, every sheet (visible and hidden), in full."""
    import io
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError(
            'openpyxl is not installed on the server '
            '(pip install openpyxl; already declared in pns_ai_mcp)'
        )
    # read_only streams; data_only gives computed values when the file has them.
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        # sheetnames lists every sheet of the workbook, hidden ones included.
        names = list(wb.sheetnames or [])
        if not names:
            return 'Excel workbook: 0 sheets (empty)'
        sheet_rows = []
        for title in names:
            ws = wb[title]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ['' if c is None else str(c) for c in row]
                while cells and cells[-1] == '':
                    cells.pop()
                if cells:
                    rows.append('\t'.join(cells))
            sheet_rows.append((title, rows))
        return format_excel_sheets(sheet_rows)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def xls_to_text(raw):
    """An old .xls (BIFF/OLE) as text, every sheet, through xlrd."""
    try:
        import xlrd
    except ImportError:
        raise ImportError(
            'xlrd is not installed on the server '
            '(pip install "xlrd>=1.2.0,<2" for .xls)'
        )
    book = xlrd.open_workbook(file_contents=raw, formatting_info=False)
    n = book.nsheets
    if n <= 0:
        return 'Excel workbook: 0 sheets (empty)'
    sheet_rows = []
    for idx in range(n):
        sh = book.sheet_by_index(idx)
        title = sh.name or ('Sheet%d' % (idx + 1))
        rows = []
        for r in range(sh.nrows):
            cells = []
            for c in range(sh.ncols):
                val = sh.cell_value(r, c)
                if val is None or val == '':
                    cells.append('')
                else:
                    cells.append(str(val))
            while cells and cells[-1] == '':
                cells.pop()
            if cells:
                rows.append('\t'.join(cells))
        sheet_rows.append((title, rows))
    return format_excel_sheets(sheet_rows)


def docx_to_text(raw):
    """The text of a .docx with no library: unzip and strip the XML tags."""
    import io
    import re
    import zipfile
    import html as _html
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        xml = z.read('word/document.xml').decode('utf-8', 'replace')
    # Paragraph and line breaks become newlines before the tags go.
    xml = re.sub(r'</w:p>', '\n', xml)
    xml = re.sub(r'<w:br\s*/?>', '\n', xml)
    text = re.sub(r'<[^>]+>', '', xml)
    return _html.unescape(text)


def _reason(exc):
    return getattr(exc, 'args', [''])[0] or type(exc).__name__


def file_text(name, mimetype, b64, ref=None):
    """The text of a data file, by its kind. ``''`` when it cannot be read."""
    if not b64:
        return ''
    try:
        raw = base64.b64decode(b64)
    except Exception:
        return ''
    lname = (name or '').lower()
    mt = (mimetype or '').lower()
    ooxml = peek_ooxml_kind(raw)
    is_ole = is_ole_cfb(raw)
    # The short .xls / .doc extensions (not .xlsx / .docx).
    is_xls_ext = lname.endswith('.xls') and not lname.endswith('.xlsx')
    is_doc_ext = lname.endswith('.doc') and not lname.endswith('.docx')
    try:
        # Excel OOXML (.xlsx / .xlsm), by extension, MIME or ZIP signature.
        if (
            lname.endswith('.xlsx')
            or lname.endswith('.xlsm')
            or 'spreadsheetml' in mt
            or ooxml == 'xlsx'
        ):
            try:
                return xlsx_to_text(raw)
            except ImportError as exc:
                return '[Excel %s: %s]' % (name, exc)
            except Exception as exc:
                _logger.warning(
                    'Chatboo file text (%s): openpyxl failed on %s: %r',
                    ref, name, exc,
                )
                return '[Excel %s: could not be read (%s)]' % (name, _reason(exc))
        # Excel BIFF (a real .xls / OLE), every sheet through xlrd.
        if is_ole and (is_xls_ext or 'ms-excel' in mt or mt == 'application/vnd.ms-excel'):
            try:
                return xls_to_text(raw)
            except ImportError as exc:
                return (
                    '[Excel %s: %s. The user can save it as .xlsx and attach it again.]'
                    % (name, exc)
                )
            except Exception as exc:
                _logger.warning(
                    'Chatboo file text (%s): xlrd failed on %s: %r',
                    ref, name, exc,
                )
                return '[Excel %s: could not be read (.xls) (%s)]' % (name, _reason(exc))
        # Word OOXML (.docx)
        if (
            lname.endswith('.docx')
            or 'wordprocessingml' in mt
            or ooxml == 'docx'
        ):
            try:
                return docx_to_text(raw)
            except Exception as exc:
                _logger.warning(
                    'Chatboo file text (%s): docx failed on %s: %r',
                    ref, name, exc,
                )
                return '[Word %s: could not be read (%s)]' % (name, _reason(exc))
        if is_ole and (is_doc_ext or 'msword' in mt or mt == 'application/msword'):
            return (
                '[Word %s: old binary .doc format. The user can open it in Word, '
                'save it as .docx and attach it again.]'
                % (name,)
            )
        # Anything else: plain text or data (txt, md, csv, tsv, log, json, ...).
        return decode_text_bytes(raw)
    except Exception:
        _logger.warning(
            'Chatboo file text (%s): could not extract text from %s (%s)',
            ref, name, mimetype, exc_info=True,
        )
        return ''
